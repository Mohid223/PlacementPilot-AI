import os
import re
import json
from pathlib import Path

from services.llm_service import ask_llm


# =========================================================
# FILE TEXT EXTRACTION
# =========================================================

def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from PDF or DOCX resume.
    """

    if not file_path:
        raise ValueError("Resume file path is required.")

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            "Resume file not found."
        )

    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    elif extension == ".docx":
        return extract_from_docx(file_path)

    else:
        raise ValueError(
            "Unsupported file format. "
            "Please upload PDF or DOCX."
        )


# =========================================================
# PDF
# =========================================================

def extract_from_pdf(file_path: str) -> str:
    """
    Extract text from PDF resume.
    """

    try:
        import PyPDF2

        text = []

        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:
                    text.append(page_text)

        return clean_text(
            "\n".join(text)
        )

    except ImportError:
        raise ImportError(
            "PyPDF2 is not installed. "
            "Run: pip install PyPDF2"
        )


# =========================================================
# DOCX
# =========================================================

def extract_from_docx(file_path: str) -> str:
    """
    Extract text from DOCX resume.
    """

    try:
        from docx import Document

        document = Document(file_path)

        text = []

        # Normal paragraphs
        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text.append(
                    paragraph.text.strip()
                )

        # Tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text.strip():
                        row_text.append(
                            cell.text.strip()
                        )

                if row_text:
                    text.append(
                        " | ".join(row_text)
                    )

        return clean_text(
            "\n".join(text)
        )

    except ImportError:
        raise ImportError(
            "python-docx is not installed. "
            "Run: pip install python-docx"
        )


# =========================================================
# CLEAN TEXT
# =========================================================

def clean_text(text: str) -> str:
    """
    Clean extracted resume text.
    """

    if not text:
        return ""

    # Remove excessive spaces
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove excessive blank lines
    text = re.sub(
        r"\n{2,}",
        "\n",
        text
    )

    return text.strip()


# =========================================================
# CLEAN LLM JSON
# =========================================================

def clean_json_response(response: str) -> str:
    """
    Remove markdown/code fences from LLM response.
    """

    if not response:
        return ""

    response = response.strip()

    response = re.sub(
        r"```json",
        "",
        response,
        flags=re.IGNORECASE
    )

    response = re.sub(
        r"```",
        "",
        response
    )

    return response.strip()


# =========================================================
# RESUME ANALYSIS
# =========================================================

def analyze_resume(resume_text: str) -> dict:
    """
    Analyze resume using PlacementPilot AI.
    """

    if not resume_text or not resume_text.strip():
        raise ValueError(
            "Resume text is required."
        )

    prompt = f"""
You are PlacementPilot AI's Resume Analyzer Agent.

Analyze the student's resume carefully.

Return ONLY valid JSON.

Do NOT use markdown.
Do NOT use ```json.
Do NOT add explanations outside JSON.

Use EXACTLY this structure:

{{
    "ats_score": 0,

    "candidate_summary": "",

    "candidate": {{
        "name": "",
        "email": "",
        "phone": "",
        "linkedin": "",
        "github": "",
        "portfolio": ""
    }},

    "technical_skills": [],

    "soft_skills": [],

    "education": [],

    "projects": [],

    "experience": [],

    "certifications": [],

    "strengths": [],

    "weak_areas": [],

    "missing_skills": [],

    "improvement_suggestions": [],

    "placement_topics": [],

    "recommended_roles": [],

    "job_match": {{
        "software_engineer": 0,
        "python_developer": 0,
        "machine_learning_engineer": 0,
        "data_science": 0
    }}
}}

IMPORTANT RULES:

1. ats_score must be between 0 and 100.

2. Do NOT invent information.

3. Only use information actually present
   in the resume.

4. If something is missing, use:
   ""
   or
   []

5. Technical skills should contain
   programming languages, frameworks,
   databases, tools and technologies.

6. Soft skills should only be included
   when explicitly mentioned or clearly
   demonstrated.

7. For projects, include:
   - project name
   - technologies
   - short description

8. For experience, include:
   - company
   - role
   - duration
   - description

9. Identify weak resume sections.

10. Identify useful missing skills for
    entry-level software engineering jobs.

11. Give practical suggestions that a
    student can actually implement.

12. placement_topics should contain topics
    the student should prepare for placements.

13. recommended_roles should contain realistic
    entry-level roles based on the resume.

14. job_match percentages must be between
    0 and 100.

15. Keep the response concise but useful.

RESUME:

{resume_text}
"""

    response = ask_llm(prompt)

    if not response:
        raise ValueError(
            "LLM returned an empty response."
        )

    response = clean_json_response(
        response
    )

    try:

        analysis = json.loads(
            response
        )

    except json.JSONDecodeError as error:

        raise ValueError(
            f"LLM returned invalid JSON: {error}"
        )

    return analysis